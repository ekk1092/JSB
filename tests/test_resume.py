import json
import os
from unittest.mock import patch, MagicMock

from docx import Document
from openai import RateLimitError

from server.tools.resume import (
    sanitize_filename,
    create_resume_docx,
    create_cover_letter_docx,
    extract_job_metadata,
    call_llm_with_retry,
    _parse_retry_after,
    tailor_resume_tool,
    _METADATA_CACHE,
    LLM_MAX_RETRIES,
)


def test_sanitize_filename():
    assert sanitize_filename("John Doe") == "John_Doe"
    assert sanitize_filename("Acme, Inc.") == "Acme_Inc"
    # Leading/trailing spaces trimmed; internal spaces become underscores
    assert sanitize_filename("  spaced  out  ") == "spaced__out"
    assert sanitize_filename("") == ""


def test_create_resume_docx():
    data = {
        "name": "Jane Smith",
        "contact": {"email": "jane@example.com", "phone": "555-1234"},
        "summary": "Experienced engineer.",
        "experience": [
            {
                "title": "Engineer",
                "company": "Acme",
                "location": "NYC",
                "dates": "2020-2023",
                "responsibilities": ["Built things", "Fixed things"],
            }
        ],
        "education": [{"degree": "BS", "school": "UNCW", "location": "NC", "graduation": "2020"}],
        "skills": ["Python", "AWS"],
    }
    path = create_resume_docx(data)
    try:
        assert os.path.exists(path)
        assert path.endswith(".docx")
        doc = Document(path)
        # Name should be present
        texts = [p.text for p in doc.paragraphs]
        assert "Jane Smith" in texts
        assert "PROFESSIONAL SUMMARY" in texts
        assert "EXPERIENCE" in texts
        assert "SKILLS" in texts
    finally:
        os.unlink(path)


def test_create_cover_letter_docx():
    data = {
        "name": "Jane Smith",
        "contact": {"email": "jane@example.com", "phone": "555-1234"},
        "date": "August 2, 2026",
        "recipient": {"name": "Hiring Manager", "company": "Acme", "address": "NYC"},
        "body_paragraphs": ["Para one.", "Para two."],
    }
    path = create_cover_letter_docx(data)
    try:
        assert os.path.exists(path)
        assert path.endswith(".docx")
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs]
        assert "Jane Smith" in texts
        assert "Acme" in texts
        assert "Sincerely," in texts
    finally:
        os.unlink(path)


def test_extract_job_metadata_caches_result():
    """extract_job_metadata should cache results per job description."""
    _METADATA_CACHE.clear()

    fake_response = type("Resp", (), {"choices": [type("C", (), {"message": type("M", (), {"content": '{"company_name": "Acme", "company_location": "NYC"}'})()})()]})()
    fake_client = type("Client", (), {"chat": type("Chat", (), {"completions": type("Comp", (), {"create": lambda self, **kw: fake_response})()})()})()

    with patch("server.tools.resume.get_llm_client", return_value=fake_client):
        meta1 = extract_job_metadata("Job description for Acme in NYC")
        meta2 = extract_job_metadata("Job description for Acme in NYC")

    assert meta1 == {"company_name": "Acme", "company_location": "NYC"}
    assert meta2 == meta1
    # Cache should have exactly one entry (second call hit cache)
    assert len(_METADATA_CACHE) == 1


# ---------------------------------------------------------------------------
# 429 RateLimitError handling tests
# ---------------------------------------------------------------------------

def _make_rate_limit_error(retry_after_s=None):
    """Create a RateLimitError for testing."""
    if retry_after_s is not None:
        msg = f"Rate limit exceeded. Please retry in {retry_after_s}s."
    else:
        msg = "Rate limit exceeded."
    return RateLimitError(message=msg, response=MagicMock(), body=None)


def test_parse_retry_after_extracts_seconds():
    """_parse_retry_after should extract the retry hint from the error message."""
    exc = _make_rate_limit_error(retry_after_s=7.5)
    assert _parse_retry_after(exc) == 7.5


def test_parse_retry_after_returns_none_when_no_hint():
    """_parse_retry_after should return None when no retry hint is present."""
    exc = _make_rate_limit_error(retry_after_s=None)
    assert _parse_retry_after(exc) is None


def test_call_llm_with_retry_succeeds_after_transient_429():
    """call_llm_with_retry should retry on RateLimitError and succeed."""
    fake_response = MagicMock()
    fake_client = MagicMock()
    # First call raises 429, second call succeeds
    fake_client.chat.completions.create.side_effect = [
        _make_rate_limit_error(retry_after_s=0.01),
        fake_response,
    ]

    with patch("server.tools.resume.time.sleep"):
        result = call_llm_with_retry(fake_client, model="test", messages=[])

    assert result is fake_response
    assert fake_client.chat.completions.create.call_count == 2


def test_call_llm_with_retry_exhausts_retries():
    """call_llm_with_retry should raise after exhausting all retries."""
    fake_client = MagicMock()
    fake_client.chat.completions.create.side_effect = _make_rate_limit_error(retry_after_s=0.01)

    with patch("server.tools.resume.time.sleep"):
        try:
            call_llm_with_retry(fake_client, model="test", messages=[])
            assert False, "Should have raised RateLimitError"
        except RateLimitError:
            pass

    assert fake_client.chat.completions.create.call_count == LLM_MAX_RETRIES


def test_tailor_resume_tool_returns_friendly_error_on_429():
    """tailor_resume_tool should return a user-friendly JSON error on persistent 429."""
    fake_client = MagicMock()
    fake_client.chat.completions.create.side_effect = _make_rate_limit_error(retry_after_s=30)

    with patch("server.tools.resume.get_llm_client", return_value=fake_client):
        with patch("server.tools.resume.time.sleep"):
            with patch(
                "server.tools.resume.extract_job_metadata",
                return_value={"company_name": None, "company_location": None},
            ):
                result = tailor_resume_tool("resume text", "job desc")

    data = json.loads(result)
    assert "error" in data
    assert "quota" in data["error"].lower()
    assert "30 seconds" in data["error"]


def test_extract_job_metadata_falls_back_on_429():
    """extract_job_metadata should return null metadata on persistent 429."""
    _METADATA_CACHE.clear()

    fake_client = MagicMock()
    fake_client.chat.completions.create.side_effect = _make_rate_limit_error(retry_after_s=0.01)

    with patch("server.tools.resume.get_llm_client", return_value=fake_client):
        with patch("server.tools.resume.time.sleep"):
            meta = extract_job_metadata("Some job description")

    assert meta == {"company_name": None, "company_location": None}